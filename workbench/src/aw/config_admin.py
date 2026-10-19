"""Local configuration patches and deterministic saved-claim replay. No policy import."""
import base64
import csv
import io
import json
from copy import deepcopy
from . import configs,evaluations
from .money import fmt


def diff(before,after,path=''):
    if isinstance(before,dict) and isinstance(after,dict):
        return [item for key in sorted(before.keys()|after.keys()) for item in diff(before.get(key),after.get(key),path+'/'+key)]
    return [] if before==after else [{'path':path,'before':before,'after':after}]


def import_patch(body,filename,encoded):
    data=base64.b64decode(encoded,validate=True)
    if len(data)>1024*1024:raise ValueError('configuration file limit is 1 MB')
    if filename.lower().endswith('.csv'):
        rows=list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
    elif filename.lower().endswith('.xlsx'):
        import zipfile
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if len(z.infolist())>100 or sum(m.file_size for m in z.infolist())>10*1024*1024:raise ValueError('XLSX expansion limit')
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            rows=[]
            for row in wb.active.iter_rows():
                if any(c.data_type=='f' for c in row):raise ValueError('spreadsheet formulas are forbidden')
                rows.append([c.value for c in row])
                if len(rows)>1001:raise ValueError('maximum 1000 configuration changes')
        finally:wb.close()
    else:raise ValueError('use CSV or XLSX')
    if not rows or rows[0]!=['path','value_json'] or len(rows)>1001:raise ValueError('header must be path,value_json; maximum1000 rows')
    out=deepcopy(body)
    for row in rows[1:]:
        if len(row)!=2 or not isinstance(row[0],str) or not row[0].startswith('/'):raise ValueError('each row requires JSON pointer and JSON value')
        keys=[k.replace('~1','/').replace('~0','~') for k in row[0][1:].split('/')]
        node=out
        for key in keys[:-1]:node=node[int(key)] if isinstance(node,list) else node[key]
        key=int(keys[-1]) if isinstance(node,list) else keys[-1]
        node[key]=json.loads(str(row[1]))
    return out


def replay(conn,tenant_id,row,claim_ids=None,store=None):
    from .engine import evaluate
    ids=claim_ids or [r['claim_id'] for r in conn.fetch_all('SELECT claim_id FROM claim WHERE tenant_id=$1 AND current_evaluation_id IS NOT NULL',tenant_id)]
    if len(ids)>500:raise ValueError('maximum500 replay claims')
    results=[]
    for cid in ids:
        current=evaluations.current(conn,tenant_id,cid)
        if not current:continue
        ei,cfg=evaluations.build_engine_input(conn,tenant_id,cid,'REPLAY')
        kind=row['artifact_kind'];candidate=row['body'];key={'POLICY_CONFIG':'policy','NME_LIST':'nme','TARIFF':'tariff','CHECKLIST':'checklist'}.get(kind)
        if key:
            if kind=='POLICY_CONFIG' and candidate['config_id']!=cfg['policy']['config_id']:continue
            if kind=='TARIFF' and candidate['hospital_id']!=ei['hospital_id']:continue
            if kind=='CHECKLIST' and candidate['workflow']!=ei['workflow']:continue
            cfg[key]=candidate
            out=evaluate(ei,cfg['policy'],cfg['nme'],cfg['tariff'],cfg['checklist'])
            after=out['recommendation']['recommended_payable']
            findings=[f['kind'] for f in out['findings']]
        elif kind=='WORKFLOW_CONFIG':
            after=fmt(current['recommended_payable_paise']);findings=[]
        else:
            # Re-extract stored pages in memory; never mutate facts during replay.
            from .extraction.pipeline import load_pages
            from .extraction.providers import classify,HeuristicProvider
            findings=[]
            for d in conn.fetch_all("SELECT document_id,classified_type FROM document WHERE tenant_id=$1 AND claim_id=$2 AND status NOT IN ('SUPERSEDED','QUARANTINED')",tenant_id,cid):
                pages=load_pages(conn,store,tenant_id,d['document_id'])
                dtype,_,_=classify(pages,candidate)
                findings.append({'document_id':d['document_id'],'before':d['classified_type'],'after':dtype,'fields':len(HeuristicProvider().extract(dtype,pages)['fields'])})
            after=fmt(current['recommended_payable_paise'])
        before=fmt(current['recommended_payable_paise'])
        results.append({'claim_id':cid,'payable_before':before,'payable_after':after,'changed':before!=after,'findings':findings,
                        'config_diff':diff(current['result'].get('config_versions',{}),{kind:configs.identity(kind,candidate)}),
                        'workflow_settings':candidate if kind=='WORKFLOW_CONFIG' else None})
    return {'claims':len(results),'changed':sum(r['changed'] for r in results),'results':results,'persisted':False}
